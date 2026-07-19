from datetime import datetime, timezone
from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

from openpyxl import load_workbook

from app.engines.documents.dto import GeneratedDocument
from app.engines.documents.equipment_excel import EquipmentExcelDocumentGenerator
from app.engines.documents.equipment_pdf import EquipmentPDFDocumentGenerator
from app.services.equipment_document_mapper import EquipmentDocumentMapper
from app.services.project_document_package_service import ProjectDocumentPackageService


def _equipment_list():
    return SimpleNamespace(
        id="equipment-list-1",
        items=[
            SimpleNamespace(
                equipment_name="Котёл",
                required_quantity=5,
                calculated_quantity=3,
                is_manual=False,
                is_removed=False,
            ),
            SimpleNamespace(
                equipment_name="Фонарь",
                required_quantity=2,
                calculated_quantity=None,
                is_manual=True,
                is_removed=False,
            ),
            SimpleNamespace(
                equipment_name="Горелка",
                required_quantity=1,
                calculated_quantity=1,
                is_manual=False,
                is_removed=True,
            ),
        ],
    )


def test_equipment_mapper_excludes_removed_items_and_marks_sources():
    dto = EquipmentDocumentMapper().to_dto(_equipment_list(), "Летний поход")

    assert dto.title == "Список оборудования"
    assert [item.equipment_name for item in dto.items] == ["Котёл", "Фонарь"]
    assert dto.items[0].source == "Изменено вручную"
    assert dto.items[1].source == "Добавлено вручную"


def test_equipment_pdf_and_excel_use_final_quantities():
    dto = EquipmentDocumentMapper().to_dto(_equipment_list(), "Летний поход")

    pdf = EquipmentPDFDocumentGenerator().generate(dto)
    assert pdf.filename == "equipment_list.pdf"
    assert pdf.content_type == "application/pdf"
    assert bytes(pdf.content).startswith(b"%PDF")

    excel = EquipmentExcelDocumentGenerator().generate(dto)
    workbook = load_workbook(BytesIO(bytes(excel.content)))
    sheet = workbook["Оборудование"]
    assert sheet["A1"].value == "Список оборудования"
    assert sheet["B2"].value == "Летний поход"
    assert [cell.value for cell in sheet[5]] == [
        "Оборудование",
        "Итого, шт.",
        "Расчёт, шт.",
        "Источник",
    ]
    assert [cell.value for cell in sheet[6]] == [
        "Котёл",
        5,
        3,
        "Изменено вручную",
    ]


class FakeDocumentService:
    def _document(self, filename: str) -> GeneratedDocument:
        return GeneratedDocument(
            filename=filename,
            content_type="application/octet-stream",
            generated_at=datetime.now(timezone.utc),
            content=filename.encode(),
        )

    def generate_consolidated_pdf(self, project):
        return self._document(f"tourhub_project_{project.id}_complete.pdf")

    def generate_consolidated_excel(self, project):
        return self._document(f"tourhub_project_{project.id}_complete.xlsx")

    def generate_purchase_pdf(self, project):
        return self._document("purchase_list.pdf")

    def generate_purchase_excel(self, project):
        return self._document("purchase_list.xlsx")

    def generate_purchase_print(self, project):
        return self._document("purchase_list.html")

    def generate_equipment_pdf(self, project):
        return self._document("equipment_list.pdf")

    def generate_equipment_excel(self, project):
        return self._document("equipment_list.xlsx")


def test_project_package_contains_complete_and_compatibility_documents():
    generated = ProjectDocumentPackageService(FakeDocumentService()).generate_package(
        SimpleNamespace(id=76)
    )
    with ZipFile(BytesIO(bytes(generated.content))) as archive:
        assert set(archive.namelist()) == {
            "tourhub_project_76_complete.pdf",
            "tourhub_project_76_complete.xlsx",
            "purchase_list.pdf",
            "purchase_list.xlsx",
            "purchase_list.html",
            "equipment_list.pdf",
            "equipment_list.xlsx",
        }
