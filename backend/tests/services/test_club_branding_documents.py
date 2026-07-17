from datetime import UTC, datetime
from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook
from PIL import Image

from app.engines.documents.branding import ClubBrandingDTO
from app.engines.documents.dto import (
    DocumentItemDTO,
    GeneratedDocument,
    PurchaseDocumentDTO,
)
from app.engines.documents.excel import ExcelDocumentGenerator
from app.engines.documents.pdf import PDFDocumentGenerator
from app.services.project_document_service import ProjectDocumentService


def _logo_bytes() -> bytes:
    buffer = BytesIO()
    Image.new("RGB", (16, 8), "white").save(buffer, format="PNG")
    return buffer.getvalue()


def _purchase_dto() -> PurchaseDocumentDTO:
    return PurchaseDocumentDTO(
        purchase_list_id="purchase-1",
        title="Закупочный лист",
        items=[DocumentItemDTO(product_name="Крупа", quantity=2, unit="кг")],
    )


def test_pdf_and_excel_accept_club_branding() -> None:
    branding = ClubBrandingDTO(
        club_name="Турклуб Север",
        logo_bytes=_logo_bytes(),
        logo_mime_type="image/png",
    )

    pdf = PDFDocumentGenerator().generate(_purchase_dto(), branding)
    assert bytes(pdf.content).startswith(b"%PDF")

    excel = ExcelDocumentGenerator().generate(_purchase_dto(), branding)
    workbook = load_workbook(BytesIO(bytes(excel.content)))
    sheet = workbook["Закупка"]
    assert workbook.properties.creator == "Турклуб Север"
    assert sheet.oddHeader.center.text == "Турклуб Север"
    assert len(sheet._images) == 1


class FakeSettingsService:
    def __init__(self, branding: ClubBrandingDTO) -> None:
        self.branding = branding
        self.calls = 0

    def to_branding(self) -> ClubBrandingDTO:
        self.calls += 1
        return self.branding


class FakeMapper:
    def to_dto(self, purchase_list) -> PurchaseDocumentDTO:
        del purchase_list
        return _purchase_dto()


class CapturingGenerator:
    def __init__(self, filename: str) -> None:
        self.filename = filename
        self.brandings: list[ClubBrandingDTO | None] = []

    def generate(self, document, branding=None) -> GeneratedDocument:
        del document
        self.brandings.append(branding)
        return GeneratedDocument(
            filename=self.filename,
            content_type="application/octet-stream",
            generated_at=datetime.now(UTC),
            content=b"document",
        )


def test_project_document_service_reuses_one_brand_snapshot() -> None:
    branding = ClubBrandingDTO(club_name="Турклуб Север")
    settings_service = FakeSettingsService(branding)
    pdf_generator = CapturingGenerator("purchase.pdf")
    excel_generator = CapturingGenerator("purchase.xlsx")
    service = ProjectDocumentService(
        document_mapper=FakeMapper(),
        pdf_generator=pdf_generator,
        excel_generator=excel_generator,
        club_settings_service=settings_service,
    )
    project = SimpleNamespace(purchase_lists=[SimpleNamespace()], meal_plans=[])

    service.generate_purchase_pdf(project)
    service.generate_purchase_excel(project)

    assert settings_service.calls == 1
    assert pdf_generator.brandings == [branding]
    assert excel_generator.brandings == [branding]
