from datetime import UTC, datetime
from io import BytesIO

from openpyxl import Workbook

from app.engines.documents.branding import ClubBrandingDTO
from app.engines.documents.branding_render import (
    apply_excel_branding,
    apply_excel_table_style,
    apply_excel_title_style,
)
from app.engines.documents.dto import GeneratedDocument
from app.engines.documents.equipment_dto import EquipmentDocumentDTO


class EquipmentExcelDocumentGenerator:
    """Generate an XLSX workbook with final and calculated equipment quantities."""

    def generate(
        self,
        document: EquipmentDocumentDTO,
        branding: ClubBrandingDTO | None = None,
    ) -> GeneratedDocument:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Оборудование"
        apply_excel_branding(workbook, sheet, branding)
        sheet.append([document.title])
        sheet.append(["Проект", document.project_name])
        sheet.append(["Идентификатор списка", document.equipment_list_id])
        sheet.append([])
        sheet.append(["Оборудование", "Итого, шт.", "Расчёт, шт.", "Источник"])

        for item in document.items:
            sheet.append(
                [
                    item.equipment_name,
                    item.required_quantity,
                    item.calculated_quantity,
                    item.source,
                ]
            )

        apply_excel_title_style(sheet, row=1, last_column=4, branding=branding)
        apply_excel_table_style(
            sheet,
            header_row=5,
            last_row=sheet.max_row,
            last_column=4,
            branding=branding,
        )
        sheet.freeze_panes = "A6"
        sheet.column_dimensions["A"].width = 34
        sheet.column_dimensions["B"].width = 14
        sheet.column_dimensions["C"].width = 16
        sheet.column_dimensions["D"].width = 24
        buffer = BytesIO()
        workbook.save(buffer)
        return GeneratedDocument(
            filename="equipment_list.xlsx",
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            generated_at=datetime.now(UTC),
            content=buffer.getvalue(),
        )
