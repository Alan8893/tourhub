from datetime import UTC, datetime
from io import BytesIO

from openpyxl import Workbook

from app.engines.documents.branding import ClubBrandingDTO
from app.engines.documents.branding_render import (
    apply_excel_branding,
    apply_excel_table_style,
    apply_excel_title_style,
)
from app.engines.documents.dto import GeneratedDocument, PurchaseDocumentDTO


class ExcelDocumentGenerator:
    """Generate Russian XLSX purchase documents."""

    def generate(
        self,
        document: PurchaseDocumentDTO,
        branding: ClubBrandingDTO | None = None,
    ) -> GeneratedDocument:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Закупка"
        apply_excel_branding(workbook, sheet, branding)
        sheet.append([document.title])
        sheet.append(["Идентификатор списка", document.purchase_list_id])
        sheet.append([])
        sheet.append(["Продукт", "Количество", "Единица", "Упаковок"])

        for item in document.items:
            sheet.append(
                [
                    item.product_name,
                    item.quantity,
                    item.unit,
                    item.packages_count,
                ]
            )

        apply_excel_title_style(sheet, row=1, last_column=4, branding=branding)
        apply_excel_table_style(
            sheet,
            header_row=4,
            last_row=sheet.max_row,
            last_column=4,
            branding=branding,
        )
        sheet.freeze_panes = "A5"
        sheet.column_dimensions["A"].width = 34
        sheet.column_dimensions["B"].width = 14
        sheet.column_dimensions["C"].width = 16
        sheet.column_dimensions["D"].width = 14
        buffer = BytesIO()
        workbook.save(buffer)

        return GeneratedDocument(
            filename="purchase_list.xlsx",
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            generated_at=datetime.now(UTC),
            content=buffer.getvalue(),
        )
